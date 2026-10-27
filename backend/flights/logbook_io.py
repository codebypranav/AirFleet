"""CSV export and import: AirFleet's own format, ForeFlight, LogTen Pro, MyFlightbook and the AirFleet import template."""
import csv
import io
import zipfile
from datetime import datetime, time, timedelta, timezone as dt_timezone

import openpyxl

from . import airports
from .models import Aircraft, Flight
from .serializers import FlightSerializer

EXPORT_FIELDS = [
    'departure_airport', 'arrival_airport', 'departure_time', 'arrival_time', 'total_time',
    'registration_number', 'aircraft_type', 'aircraft_condition', 'distance',
    'pic_time', 'sic_time', 'dual_received_time', 'night_time', 'instrument_time', 'simulated_instrument_time',
    'day_landings', 'night_landings', 'approaches', 'cross_country', 'is_simulator',
    'departure_gate', 'arrival_gate', 'flight_plan', 'weather_conditions', 'notes',
    'instructor_name', 'instructor_certificate', 'instructor_signed_at', 'signature_status',
]
# From the flight's latest instructor signature. Informational: a signature can't be imported.
SIGNATURE_FIELDS = {
    'instructor_name': 'instructor_name',
    'instructor_certificate': 'certificate_number',
    'instructor_signed_at': 'signed_at',
    'signature_status': 'status',
}
# Read back on import; aircraft_type and the signature columns are informational only.
IMPORT_FIELDS = [f for f in EXPORT_FIELDS if f not in ('aircraft_type', *SIGNATURE_FIELDS)]

FOREFLIGHT_CLASSES = {
    'airplane_single_engine_land': 'SEL',
    'airplane_multi_engine_land': 'MEL',
    'airplane_single_engine_sea': 'SES',
    'airplane_multi_engine_sea': 'MES',
    'rotorcraft_helicopter': 'HELICOPTER',
    'glider': 'GLIDER',
}
MYFLIGHTBOOK_CLASSES = {'asel': 'SEL', 'amel': 'MEL', 'ases': 'SES', 'ames': 'MES', 'helicopter': 'HELICOPTER', 'glider': 'GLIDER'}

TEMPLATE = {
    'date': 'Date', 'departure_airport': 'From', 'arrival_airport': 'To', 'flight_plan': 'Route',
    'registration_number': 'Aircraft', 'type_code': 'Type', 'takeoff': 'Time Off', 'total_time': 'Total',
    'pic_time': 'PIC', 'sic_time': 'SIC', 'dual_received_time': 'Dual Received', 'night_time': 'Night',
    'instrument_time': 'Actual Instrument', 'simulated_instrument_time': 'Simulated Instrument',
    'cross_country': 'Cross Country', 'day_landings': 'Day Landings', 'night_landings': 'Night Landings',
    'approaches': 'Approaches', 'is_simulator': 'Simulator', 'notes': 'Remarks',
}
MAX_SHEET_BYTES = 50 * 1024 * 1024


def export_csv(flights):
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=EXPORT_FIELDS, extrasaction='ignore')
    writer.writeheader()
    for flight in flights.select_related('aircraft').prefetch_related('signatures'):
        row = FlightSerializer(flight).data
        signature = row.get('signature') or {}
        row.update({column: signature.get(key, '') for column, key in SIGNATURE_FIELDS.items()})
        writer.writerow({field: row.get(field, '') for field in EXPORT_FIELDS})
    return out.getvalue()


class LogbookImportError(Exception):
    pass


def _hours(value):
    value = (value or '').strip()
    try:
        if ':' in value:
            hours, minutes = value.split(':')[:2]
            return timedelta(hours=int(hours), minutes=int(minutes))
        return timedelta(hours=float(value)) if value else timedelta(0)
    except ValueError:
        return timedelta(0)


def _int(value):
    try:
        return int(float(value)) if value not in (None, '') else 0
    except ValueError:
        return 0


def _text(row, column):
    return (row.get(column) or '').strip()


def _clock(value):
    """Clock times are UTC and look like 1435, 14:35 or 14:35:00, alone or after a date."""
    parts = (value or '').strip().rstrip('Z').split()
    digits = ''.join(parts[-1].split(':')[:2]) if parts else ''
    if len(digits) in (3, 4) and digits.isdigit():
        digits = digits.zfill(4)
        return time(int(digits[:2]) % 24, int(digits[2:]) % 60)
    return None


def _day(value):
    parts = (value or '').split()
    for pattern in ('%Y-%m-%d', '%m/%d/%Y', '%m/%d/%y'):
        try:
            return datetime.strptime(parts[0] if parts else '', pattern).date()
        except ValueError:
            pass
    return None


def _departure(day, *values):
    for value in values:
        start = _clock(value)
        if start:
            return datetime.combine(_day(value) or day, start, tzinfo=dt_timezone.utc)
    return datetime.combine(day, time(12, 0), tzinfo=dt_timezone.utc)


def _leg(departure, total, **fields):
    return {
        'departure_time': departure,
        'arrival_time': departure + total,
        'total_time': total,
        **{k: min(v, total) if isinstance(v, timedelta) else v for k, v in fields.items()},
    }


def _problem(date_text, day, total):
    if not date_text or total <= timedelta(0):
        return 'Missing date or total time'
    if not day:
        return f'Unrecognised date {date_text!r}'
    return None


def _yes(value):
    return value.lower() in ('y', 'yes', 'true', 'x') or _hours(value) > timedelta(0)


def _airport(code):
    code = code.strip().upper()
    if len(code) != 3 or airports.get(code):
        return code
    readings = {airports.from_iata(code), 'K' + code if airports.get('K' + code) else None} - {None}
    if len(readings) > 1:
        raise ValueError(f"{code} could be {' or '.join(sorted(readings))}. Use the 4-letter ICAO code.")
    return readings.pop() if readings else code


def _tables(text):
    """Split a ForeFlight export into its named tables ("Aircraft Table", "Flights Table")."""
    tables, name, rows = {}, None, []
    for row in csv.reader(io.StringIO(text)):
        first = row[0].strip() if row else ''
        if first.endswith(' Table'):
            if name:
                tables[name] = rows
            name, rows = first, []
        elif name is not None:
            rows.append(row)
    if name:
        tables[name] = rows
    parsed = {}
    for name, rows in tables.items():
        rows = [r for r in rows if any(cell.strip() for cell in r)]
        if not rows:
            continue
        header = [h.strip() for h in rows[0]]
        parsed[name] = [dict(zip(header, r)) for r in rows[1:]]
    return parsed


def _approach_count(row):
    total = 0
    for i in range(1, 7):
        value = (row.get(f'Approach{i}') or '').strip()
        if value:
            first = value.split(';')[0]
            total += int(first) if first.isdigit() else 1
    return total


def parse_foreflight(text):
    tables = _tables(text)
    if 'Flights Table' not in tables:
        raise LogbookImportError("This doesn't look like a ForeFlight logbook export (no Flights Table found).")

    aircraft = {}
    for row in tables.get('Aircraft Table', []):
        reg = (row.get('AircraftID') or '').strip().upper()
        if reg:
            aircraft[reg] = {
                'type_code': (row.get('TypeCode') or '').strip().upper()[:4],
                'make_model': ' '.join(filter(None, [(row.get('Make') or '').strip(), (row.get('Model') or '').strip()]))[:60],
                'aircraft_class': FOREFLIGHT_CLASSES.get((row.get('Class') or '').strip().lower(), 'OTHER'),
            }

    flights = []
    for row in tables['Flights Table']:
        total = _hours(row.get('TotalTime'))
        if not row.get('Date') or total <= timedelta(0):
            flights.append({'_error': 'Missing date or total time', '_raw': row})
            continue
        try:
            day = datetime.strptime(row['Date'].strip(), '%Y-%m-%d').date()
        except ValueError:
            flights.append({'_error': f"Unrecognised date {row['Date']!r}", '_raw': row})
            continue
        night_landings = _int(row.get('NightLandingsFullStop'))
        all_landings = _int(row.get('AllLandings'))
        flights.append(_leg(
            _departure(day, row.get('TimeOff'), row.get('TimeOut')), total,
            departure_airport=(row.get('From') or '').strip().upper(),
            arrival_airport=(row.get('To') or row.get('From') or '').strip().upper(),
            registration_number=(row.get('AircraftID') or '').strip().upper() or 'UNKNOWN',
            distance=_int(row.get('Distance')),
            pic_time=_hours(row.get('PIC')),
            sic_time=_hours(row.get('SIC')),
            dual_received_time=_hours(row.get('DualReceived')),
            night_time=_hours(row.get('Night')),
            instrument_time=_hours(row.get('ActualInstrument')),
            simulated_instrument_time=_hours(row.get('SimulatedInstrument')),
            day_landings=max(all_landings - night_landings, _int(row.get('DayLandingsFullStop'))),
            night_landings=night_landings,
            approaches=_approach_count(row),
            cross_country=_hours(row.get('CrossCountry')) > timedelta(0),
            is_simulator=_hours(row.get('SimulatedFlight')) > timedelta(0),
            flight_plan=(row.get('Route') or '').strip(),
            notes=(row.get('PilotComments') or '').strip(),
        ))
    return flights, aircraft


def parse_logten(rows):
    flights, aircraft = [], {}
    for row in rows:
        date_text = _text(row, 'flight_flightDate')
        day = _day(date_text)
        total = _hours(row.get('flight_totalTime')) or _hours(row.get('flight_simulator'))
        problem = _problem(date_text, day, total)
        if problem:
            flights.append({'_error': problem, '_raw': row})
            continue
        registration = _text(row, 'aircraft_aircraftID').upper() or 'UNKNOWN'
        origin = _text(row, 'flight_from').upper()
        flights.append(_leg(
            _departure(day, row.get('flight_takeoffTime'), row.get('flight_actualDepartureTime')), total,
            departure_airport=origin,
            arrival_airport=_text(row, 'flight_to').upper() or origin,
            registration_number=registration,
            distance=_int(row.get('flight_distance')),
            pic_time=_hours(row.get('flight_pic')),
            sic_time=_hours(row.get('flight_sic')),
            dual_received_time=_hours(row.get('flight_dualReceived')),
            night_time=_hours(row.get('flight_night')),
            instrument_time=_hours(row.get('flight_actualInstrument')),
            simulated_instrument_time=_hours(row.get('flight_simulatedInstrument')),
            day_landings=_int(row.get('flight_dayLandings')),
            night_landings=_int(row.get('flight_nightLandings')),
            approaches=_int(row.get('flight_totalApproaches')),
            cross_country=_hours(row.get('flight_crossCountry')) > timedelta(0),
            is_simulator=_hours(row.get('flight_simulator')) > timedelta(0),
            flight_plan=_text(row, 'flight_route'),
            notes=_text(row, 'flight_remarks'),
        ))
        aircraft.setdefault(registration, {
            'type_code': _text(row, 'aircraftType_type').upper()[:4],
            'make_model': ' '.join(filter(None, [_text(row, 'aircraftType_make'), _text(row, 'aircraftType_model')]))[:60],
            'aircraft_class': 'OTHER',
        })
    return flights, aircraft


def parse_myflightbook(rows):
    flights, aircraft = [], {}
    for row in rows:
        date_text = _text(row, 'Date')
        day = _day(date_text)
        total = _hours(row.get('Total Flight Time')) or _hours(row.get('Ground Simulator'))
        problem = _problem(date_text, day, total)
        if problem:
            flights.append({'_error': problem, '_raw': row})
            continue
        registration = _text(row, 'Tail Number').upper() or 'UNKNOWN'
        stops = _text(row, 'Route').upper().split() or ['']
        night_landings = _int(row.get('FS Night Landings'))
        flights.append(_leg(
            _departure(day, row.get('Flight Start'), row.get('Engine Start')), total,
            departure_airport=stops[0],
            arrival_airport=stops[-1],
            registration_number=registration,
            pic_time=_hours(row.get('PIC')),
            sic_time=_hours(row.get('SIC')),
            dual_received_time=_hours(row.get('Dual Received')),
            night_time=_hours(row.get('Night')),
            instrument_time=_hours(row.get('IMC')),
            simulated_instrument_time=_hours(row.get('Simulated Instrument')),
            day_landings=max(_int(row.get('Landings')) - night_landings, _int(row.get('FS Day Landings'))),
            night_landings=night_landings,
            approaches=_int(row.get('Approaches')),
            cross_country=_hours(row.get('X-Country')) > timedelta(0),
            is_simulator=_hours(row.get('Ground Simulator')) > timedelta(0),
            flight_plan=_text(row, 'Route'),
            notes=_text(row, 'Comments'),
        ))
        aircraft.setdefault(registration, {
            'type_code': _text(row, 'ICAO Model').upper()[:4],
            'make_model': _text(row, 'Model')[:60],
            'aircraft_class': MYFLIGHTBOOK_CLASSES.get(_text(row, 'Category/Class').lower(), 'OTHER'),
        })
    return flights, aircraft


def parse_template(rows):
    flights, aircraft = [], {}
    for row in rows:
        cell = {key: _text(row, column) for key, column in TEMPLATE.items()}
        day, total = _day(cell['date']), _hours(cell['total_time'])
        problem = _problem(cell['date'], day, total)
        if problem:
            flights.append({'_error': problem, '_raw': row})
            continue
        registration = cell['registration_number'].upper() or 'UNKNOWN'
        origin = cell['departure_airport'].upper()
        flights.append(_leg(
            _departure(day, cell['takeoff']), total,
            departure_airport=origin,
            arrival_airport=cell['arrival_airport'].upper() or origin,
            registration_number=registration,
            **{k: _hours(cell[k]) for k in ('pic_time', 'sic_time', 'dual_received_time', 'night_time', 'instrument_time', 'simulated_instrument_time')},
            **{k: _int(cell[k]) for k in ('day_landings', 'night_landings', 'approaches')},
            cross_country=_yes(cell['cross_country']),
            is_simulator=_yes(cell['is_simulator']),
            flight_plan=cell['flight_plan'],
            notes=cell['notes'],
        ))
        if cell['type_code']:
            aircraft.setdefault(registration, {'type_code': cell['type_code'].upper()[:4], 'make_model': '', 'aircraft_class': 'OTHER'})
    return flights, aircraft


def parse_airfleet(header, rows):
    missing = {'departure_airport', 'arrival_airport', 'departure_time', 'arrival_time', 'registration_number'} - header
    if missing:
        raise LogbookImportError(f"Missing columns: {', '.join(sorted(missing))}. Export a CSV from AirFleet to see the format.")
    return [{k: v for k, v in row.items() if k in IMPORT_FIELDS and v not in (None, '')} for row in rows], {}


def template_xlsx():
    book = openpyxl.Workbook()
    book.active.title = 'Flights'
    book.active.append(list(TEMPLATE.values()))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _cell(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M')
    if isinstance(value, time):
        return value.strftime('%H:%M')
    if isinstance(value, timedelta):
        minutes = round(value.total_seconds() / 60)
        return f'{minutes // 60}:{minutes % 60:02d}'
    return str(value)


def _sheet(data):
    unreadable = "Couldn't read this spreadsheet. Save it as .xlsx or CSV and try again."
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            size = sum(member.file_size for member in archive.infolist())
    except zipfile.BadZipFile as e:
        raise LogbookImportError(unreadable) from e
    if size > MAX_SHEET_BYTES:
        raise LogbookImportError('This spreadsheet is too large to import. Split it into smaller files.')
    try:
        cells = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True).active.iter_rows(values_only=True)
        header = [_cell(value).strip() for value in next(cells, ())]
        rows = [dict(zip(header, map(_cell, values))) for values in cells if any(v not in (None, '') for v in values)]
    except Exception as e:
        raise LogbookImportError(unreadable) from e
    return set(header), rows


def _table(text):
    first_line = text.split('\n', 1)[0]
    reader = csv.DictReader(io.StringIO(text), delimiter='\t' if '\t' in first_line else ',')
    reader.fieldnames = [name.strip() for name in reader.fieldnames or []]
    return set(reader.fieldnames), list(reader)


def detect_and_parse(data):
    if data[:2] == b'PK':
        header, rows = _sheet(data)
    else:
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('latin-1')
        head = text[:4000]
        first_line = head.splitlines()[0] if head else ''
        if 'Flights Table' in head or 'ForeFlight' in first_line:
            return 'foreflight', *parse_foreflight(text)
        header, rows = _table(text)
    if {'flight_flightDate', 'aircraft_aircraftID'} <= header:
        return 'logten', *parse_logten(rows)
    if {'Date', 'Tail Number', 'Total Flight Time'} <= header:
        return 'myflightbook', *parse_myflightbook(rows)
    if {TEMPLATE['date'], TEMPLATE['departure_airport'], TEMPLATE['registration_number'], TEMPLATE['total_time']} <= header:
        return 'template', *parse_template(rows)
    if 'departure_airport' in header:
        return 'airfleet', *parse_airfleet(header, rows)
    raise LogbookImportError(
        "Couldn't recognise this file. AirFleet reads ForeFlight and MyFlightbook CSV exports, LogTen Pro tab-delimited "
        "exports, the AirFleet import template (.xlsx or CSV) and its own CSV export."
    )


def import_flights(upload, request):
    """Create every valid row, skipping duplicates. Returns a summary with per-row problems."""
    source, rows, aircraft_info = detect_and_parse(upload)
    user = request.user
    existing = set(Flight.objects.filter(user=user).values_list('departure_time', 'departure_airport', 'arrival_airport'))

    created, duplicates, errors = 0, 0, []
    for line, row in enumerate(rows, start=1):
        if '_error' in row:
            errors.append({'row': line, 'errors': row['_error']})
            continue
        try:
            row.update({key: _airport(row[key]) for key in ('departure_airport', 'arrival_airport') if row.get(key)})
        except ValueError as e:
            errors.append({'row': line, 'errors': str(e)})
            continue
        serializer = FlightSerializer(data=row, context={'request': request, 'historical': True})
        if not serializer.is_valid():
            errors.append({'row': line, 'errors': serializer.errors})
            continue
        data = serializer.validated_data
        key = (data['departure_time'], data['departure_airport'], data['arrival_airport'])
        if key in existing:
            duplicates += 1
            continue
        serializer.save(user=user)
        existing.add(key)
        created += 1

    for reg, info in aircraft_info.items():
        plane = Aircraft.objects.filter(user=user, registration=reg).first()
        if plane:
            # Fill in what the pilot hasn't set; SEL is the model default, so treat it as unset.
            changed = {k: v for k, v in info.items() if k != 'aircraft_class' and v and not getattr(plane, k)}
            if plane.aircraft_class == 'SEL' and info['aircraft_class'] not in ('SEL', 'OTHER'):
                changed['aircraft_class'] = info['aircraft_class']
            for k, v in changed.items():
                setattr(plane, k, v)
            if changed:
                plane.save()

    return {'source': source, 'created': created, 'duplicates': duplicates, 'skipped': len(errors), 'errors': errors[:50]}
